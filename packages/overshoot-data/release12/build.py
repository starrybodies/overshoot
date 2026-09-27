"""Normalize WAW 3.0 workbooks and expose retained IRP histories by place.
Run with Python + openpyxl from the repository root. No network calls.
"""
import collections, hashlib, json, math, pathlib, re
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[3]
HERE=pathlib.Path(__file__).parent
OUT=ROOT/'public/data/v12'
OUT.mkdir(exist_ok=True)
(OUT/'waste').mkdir(exist_ok=True)
(OUT/'history').mkdir(exist_ok=True)
def dump(path,value): path.write_text(json.dumps(value,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\n')
def numeric(v): return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)
fields={}
def field(key,label,group,unit,description): fields[key]={'label':label,'group':group,'unit':unit,'description':description}
field('tonnes','Waste generated','generation','tonnes/year','Municipal waste: households and similar waste from businesses and institutions. Measurement boundaries differ by place.')
field('perPerson','Waste per person','generation','kg/person/day','The source’s municipal waste quantity relative to population. It is an area average, not a household measurement.')
for k,l in [('food_organic_waste','Food & organic material'),('glass','Glass'),('metal','Metal'),('paper_cardboard','Paper & cardboard'),('plastic','Plastic'),('rubber_leather','Rubber & leather'),('wood','Wood'),('yard_garden_green_waste','Garden & green waste'),('textile_waste','Textiles'),('weee','Electrical equipment'),('hazardous','Hazardous material'),('diapers_and_hygiene','Diapers & hygiene'),('other','Other material')]:
 field('composition_'+k,l,'composition','% of municipal waste','Share by weight of municipal waste, as categorized in the source. Definitions and sampling years may vary.')
for k,l,desc in [
 ('open_dumpsite','Open dumping','Disposal with no or limited controls.'),
 ('controlled_landfill','Controlled landfill','Land disposal with basic operational controls.'),
 ('sanitary_landfill_landfill_gas_system','Sanitary landfill','Land disposal with improved or full controls under the source classification.'),
 ('landfill_unspecified','Landfill · type unspecified','The source does not establish the level of landfill control.'),
 ('anaerobic_digestion','Anaerobic digestion','Biological treatment without oxygen, producing gas and a digestate.'),
 ('compost','Composting','Biological processing of organic material.'),
 ('recycling','Sent for recycling','The World Bank defines this as waste collected for recycling. It does not establish the quantity of usable material recovered after processing losses.'),
 ('incineration','Incineration','Controlled thermal treatment, with or without energy recovery. Open burning is excluded.'),
 ('mbt','Mechanical & biological treatment','Combined sorting and biological processing. This is a treatment stage rather than proof of a final recycled product.'),
 ('rdf','Fuel preparation','Processing into refuse-derived fuel. Subsequent burning is a separate step.'),
 ('other','Other treatment','Treatment not otherwise specified in the source.'),
 ('unaccounted_for','Unaccounted for','A gap explicitly reported by the source; OVERSHOOT does not calculate or redistribute it.')]:
 field('treatment_'+k,l,'treatment','% of municipal waste',desc)
field('uncollected','Uncollected','treatment','% of municipal waste','Municipal waste outside the recorded collection system. Its destination is not inferred.')
for k,l,unit in [('households','Households with collection','% of households'),('population','Population with collection','% of population'),('waste','Waste collected','% of municipal waste')]:field('collection_'+k,l,'collection',unit,'Collection coverage uses the stated denominator. Population, household and waste-weight percentages are different measures.')
for k,l in [('agricultural','Agricultural waste'),('construction_and_demolition','Construction & demolition waste'),('e','Electronic waste'),('hazardous','Hazardous waste'),('industrial','Industrial waste, excluding mining'),('mining','Mining waste'),('medical','Medical waste')]:field('other_'+k,l,'other','tonnes/year','A separate non-municipal waste measure. It is not added to municipal waste or overlapping categories.')
for y in [2022,2030,2040,2050]:field('projection_'+str(y),str(y)+' modeled waste','projection','tonnes/year','World Bank modeled municipal waste generation for this year. These are projections, not observed quantities. The 2022 estimate is the comparable baseline.')
field('distance','Distance to main treatment or disposal facility','local','km','Reported distance from city centre to the main land disposal or recovery facility. It is not a traced shipment or the distance traveled by all waste.')
mapping={'msw_total_msw_generated_tonnes_year':'tonnes','msw_total_msw_generated_tons_per_year':'tonnes','msw_total_msw_generated_kg_per_cap_per_day':'perPerson','waste_uncollected_percent':'uncollected','transport_distance_from_city_center_to_main_land_disposal_or_recovery_facility_km':'distance'}
for k in fields:
 if k.startswith('composition_'):mapping['composition_msw_'+k.removeprefix('composition_')+'_percent']=k
 if k.startswith('treatment_'):mapping['waste_treatment_'+k.removeprefix('treatment_')+'_percent']=k
 if k.startswith('collection_'):mapping['waste_collection_coverage_total_percent_of_'+k.removeprefix('collection_')]=k
 if k.startswith('projection_'):mapping['msw_total_msw_generated_tonnes_year_projected_'+k.removeprefix('projection_')]=k
other={'agricultural':'agricultural_waste','construction_and_demolition':'construction_and_demolition_waste','e':'e_waste','hazardous':'hazardous_waste','industrial':'industrial_waste_generated_excluding_mining','mining':'mining_waste','medical':'medical_waste'}
for k,v in other.items():mapping['non_msw_'+v+('_tonnes_year' if k=='industrial' else '_generated_tonnes_year')]='other_'+k
# City numeric values use the same fractional-percent storage as the country workbook.
# Valparaíso population coverage contains 94.6 rather than 0.946. Retain the
# original cell and withhold this out-of-range observation rather than guessing.
summary=[];counts=collections.Counter(); anomalies=[]
for level,filename in [('country','waw3-countries.xlsx'),('city','waw3-cities.xlsx')]:
 workbook=load_workbook(HERE/'raw'/filename,read_only=True,data_only=True)
 codebook=workbook['Codebook']; cbrows=codebook.iter_rows(max_col=15,values_only=True); cbheader=next(cbrows);refs=collections.defaultdict(list);empty=0
 for rownum,row in enumerate(cbrows,2):
  if not any(row):
   empty+=1
   if empty>100:break
   continue
  empty=0;d=dict(zip(cbheader,row));country=str(d.get('iso3c') or '').strip();city=str(d.get('city_code') or '').strip();m=str(d.get('measurement') or '').strip()
  ref={'workbook':filename,'sheet':'Codebook','row':rownum,'date':d.get('date_of_measurement'),'source':d.get('source/reference',d.get('source')),'page':d.get('source/reference_page_figure',d.get('page_figure')),'url':d.get('source/reference_url',d.get('weblink')),'method':d.get('method_of_measurement'),'point':d.get('point_of_measurement',d.get('point_of_measuremnet')),'notes':d.get('notes'),'collectionNotes':d.get('additional_explanation_for_method_of_data_collection',d.get('additional_explanation_for_methof_of_data_collection'))}
  ref={k:v for k,v in ref.items() if v is not None and v!=''}
  refs[(country,city if level=='city' else '',m)].append(ref)
 sheet=workbook.worksheets[1]; rows=sheet.iter_rows(values_only=True)
 if level=='country':next(rows)
 keys=next(rows)
 for rownum,row in enumerate(rows,3 if level=='country' else 2):
  d=dict(zip(keys,row));country=str(d.get('iso3c') or '').strip()
  if not re.fullmatch('[A-Z]{3}',country):continue
  city=str(d.get('city_code') or '').strip() if level=='city' else ''
  id=country+'-'+city if city else country
  record={'id':id,'country':country,'level':level,'name':d.get('city_name') if city else d.get('country_name'),'countryName':d.get('country_name'),'region':d.get('region_id'),'year':d.get('msw_total_msw_generation_year'),'point':d.get('msw_total_msw_generated_point_of_measurement'),'sourceId':'worldbank-waw3-2026','workbook':filename,'sheet':sheet.title,'row':rownum,'observations':{}}
  for sourceKey,key in mapping.items():
   if sourceKey not in d:continue
   raw=d[sourceKey]
   if raw is None or raw=='':continue
   percent=fields[key]['unit'].startswith('%')
   value=raw*100 if numeric(raw) and percent else raw if numeric(raw) else None
   issue=None
   if value is not None and (value<0 or percent and value>100):
    issue='Source cell falls outside the expected range. No corrected value is inferred.';anomalies.append({'record':id,'field':key,'rawValue':raw});value=None
   reference=refs.get((country,city,sourceKey),[])
   # Preserve the city workbook's tons/tonnes spelling when matching its codebook.
   if not reference and 'tons_per_year' in sourceKey:reference=refs.get((country,city,sourceKey.replace('tons_per_year','tonnes_year')),[])
   obs={'value':value,'rawValue':raw,'sourceField':sourceKey,'unit':fields[key]['unit'],'references':reference}
   if issue:obs['issue']=issue
   record['observations'][key]=obs
   counts['observations']+=value is not None
   counts['referencedObservations']+=bool(reference)
  assert record['name'] and id not in [r['id'] for r in summary]
  dump(OUT/'waste'/f'{id}.json',record)
  entry={k:v for k,v in record.items() if k not in ['observations','point','workbook','sheet','row','sourceId']}
  for k in ['tonnes','perPerson','projection_2022','treatment_recycling','treatment_open_dumpsite']:
   o=record['observations'].get(k)
   entry[k]=o['value'] if o else None
   if o and o['references']:entry[k+'Dates']=list(dict.fromkeys(str(x['date']) for x in o['references'] if x.get('date')))
  entry['compositionCount']=sum(k.startswith('composition_') and o['value'] is not None for k,o in record['observations'].items())
  entry['treatmentCount']=sum((k.startswith('treatment_') or k=='uncollected') and o['value'] is not None for k,o in record['observations'].items())
  summary.append(entry);counts[level+'Records']+=1
assert counts['countryRecords']==217 and counts['cityRecords']==262
catalog={'sourceId':'worldbank-waw3-2026','title':'What a Waste 3.0','publisher':'World Bank','edition':2026,'retrievedAt':'2026-09-21','license':'CC BY 4.0','url':'https://datacatalog.worldbank.org/search/dataset/0039597/what-a-waste-global-database','citation':'Cook, Ed, Kremena Ionkova, Perinaz Bhada-Tata, Sonakshi Yadav, and Frank van Woerden. 2026. What a Waste 3.0: Global Snapshot of Solid Waste Management Toward Circularity until 2050. World Bank. Country and city datasets.','counts':dict(counts),'fields':fields,'records':summary,'anomalies':anomalies}
dump(OUT/'waste-index.json',catalog)
# Country-specific histories avoid loading all 55 global annual files in a browser.
histories=collections.defaultdict(list)
historyFields=['year','extraction','imports','exports','domesticConsumption','footprint','population','extractionPerCapita','domesticConsumptionPerCapita','footprintPerCapita','estimated']
for file in sorted((ROOT/'public/data/v2/resources').glob('*.json')):
 for row in json.loads(file.read_text())['rows']:histories[row['country']].append({k:row.get(k) for k in historyFields})
for country,rows in histories.items():dump(OUT/'history'/f'{country}.json',{'sourceId':'irp-resource-accounts-2026','country':country,'rows':sorted(rows,key=lambda r:r['year'])})
report={'waste':dict(counts),'anomalies':anomalies,'historyPlaces':len(histories),'historyRecords':sum(map(len,histories.values()))}
dump(HERE/'validation-summary.json',report)
print(json.dumps(report,indent=2))
