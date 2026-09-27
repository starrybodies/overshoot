"""Check published records against retained source cells and previous IRP data.

Independent of build.py: validate every emitted waste cell and reference location,
all annual history values, source checksums and exact-year production records.
Includes source anomalies that must never become plausible-looking corrections.
"""
import collections, csv, hashlib, json, math, pathlib, re
from openpyxl import load_workbook

HERE=pathlib.Path(__file__).parent
ROOT=HERE.parents[2]
PUBLIC=ROOT/'public/data/v12'
read=lambda p:json.loads(p.read_text())
manifest=read(HERE/'manifest.json')
index=read(PUBLIC/'waste-index.json')
assert len(index['records'])==479
assert collections.Counter(r['level'] for r in index['records'])=={'country':217,'city':262}
assert len({r['id'] for r in index['records']})==479
source_tables={};source_notes={}
for source in manifest['files']:
    path=HERE/'raw'/source['file']
    assert hashlib.sha256(path.read_bytes()).hexdigest()==source['sha256']
    wb=load_workbook(path,read_only=True,data_only=True)
    rows=wb.worksheets[1].iter_rows(values_only=True)
    country=source['file']=='waw3-countries.xlsx'
    if country:next(rows)
    fields=next(rows)
    source_tables[source['file']]={i:dict(zip(fields,row)) for i,row in enumerate(rows,3 if country else 2)}
    rows=wb['Codebook'].iter_rows(max_col=15,values_only=True)
    header=next(rows);notes={};empty=0
    for i,row in enumerate(rows,2):
        if not any(row):
            empty+=1
            if empty>100:break
            continue
        empty=0;notes[i]=dict(zip(header,row))
    source_notes[source['file']]=notes
    wb.close()

numeric=zero=references=withheld=0
for entry in index['records']:
    assert re.fullmatch(r'[A-Z]{3}(?:-[A-Za-z0-9_-]{1,20})?',entry['id']),entry['id']
    r=read(PUBLIC/'waste'/f"{entry['id']}.json")
    original=source_tables[r['workbook']][r['row']]
    assert r['country']==original['iso3c']
    assert r['year']==original['msw_total_msw_generation_year']
    assert r['name']==original['city_name' if r['level']=='city' else 'country_name']
    if r['level']=='city':assert r['id']==r['country']+'-'+original['city_code']
    for key,o in r['observations'].items():
        assert o['rawValue']==original[o['sourceField']],(entry['id'],key)
        assert o['unit']==index['fields'][key]['unit']
        if o['value'] is not None:
            numeric+=1
            assert math.isfinite(o['value']) and o['value']>=0
            expected=o['rawValue']*(100 if o['unit'].startswith('%') else 1)
            assert o['value']==expected
            if o['unit'].startswith('%'):assert o['value']<=100
            if o['value']==0:zero+=1
        else:
            withheld+=1
            assert o.get('issue') or isinstance(o['rawValue'],str)
        for ref in o['references']:
            original_ref=source_notes[ref['workbook']][ref['row']]
            assert original_ref['iso3c']==r['country']
            if r['level']=='city':assert original_ref['city_code']==original['city_code']
            assert str(original_ref['measurement']).replace('tons_per_year','tonnes_year')==o['sourceField'].replace('tons_per_year','tonnes_year')
            assert ref.get('date')==original_ref.get('date_of_measurement')
            references+=1
    for key in ['tonnes','perPerson','projection_2022','treatment_recycling','treatment_open_dumpsite']:
        assert entry[key]==r['observations'].get(key,{}).get('value')
assert numeric==index['counts']['observations']==6717
assert zero>0 and withheld>=1
can=read(PUBLIC/'waste/CAN.json')
assert can['year']==2018
assert can['observations']['tonnes']['value']==30452356
assert 2016 in [r['date'] for r in can['observations']['tonnes']['references']]
bad=read(PUBLIC/'waste/CHL-valp.json')['observations']['collection_population']
assert bad['rawValue']==94.6 and bad['value'] is None and bad['issue']
assert all(index['fields']['projection_'+str(y)]['group']=='projection' for y in [2022,2030,2040,2050])

history={p.stem:read(p) for p in (PUBLIC/'history').glob('*.json')}
count=negative=missing=0
for p in (ROOT/'public/data/v2/resources').glob('*.json'):
    for raw in read(p)['rows']:
        same=next(r for r in history[raw['country']]['rows'] if r['year']==raw['year'])
        assert all(value==raw.get(key) for key,value in same.items())
        negative+=sum(isinstance(v,(int,float)) and v<0 for v in same.values())
        missing+=sum(v is None for v in same.values())
        count+=1
assert count==12705 and len(history)==231
assert negative>0 and missing>0
assert all([r['year'] for r in h['rows']]==list(range(1970,2025)) for h in history.values())

forest_raw=list(csv.DictReader((HERE/'raw/faostat-forestry-2024-subset.csv').open()))
countries={c['numeric']:c['id'] for c in read(ROOT/'packages/overshoot-data/expanded/resources/raw/country-contract.json')}
forest={(('WORLD' if r['Area']=='World' else countries[r['Area Code (M49)'].lstrip("'")]),r['Item Code']):r for r in forest_raw}
unique=set()
for family in ['wood','paper']:
    dataset=read(PUBLIC/'production'/f'{family}.json')
    for row in dataset['rows']:
        unique.add((row['country'],row['item_code']))
        original=forest[(row['country'],row['item_code'])]
        assert row['value']==(float(original['Value']) if original['Value'] else None)
        assert row['year']==int(original['Year'])==2024
        assert row['source_flag']==original['Flag']
        assert row['unit']=={'m3':'m³','t':'tonnes'}[original['Unit']]
        assert row['original_unit']==original['Unit']
assert len(unique)==1388
crop=read(PUBLIC/'production/food.json')
original_crop=read(ROOT/'public/data/v2/agriculture.json')
assert len(crop['rows'])==len(original_crop['rows'])==963
for row,raw in zip(crop['rows'],original_crop['rows']):
    assert row['value']==raw['tonnes'] and row['country']==raw['country'] and row['item_code']==raw['item_code']
    assert row['source_flag']==raw['source_flag'] and row['year']==raw['year']==2024
    assert row['unit']=='tonnes'
print(json.dumps({'result':'PASS','waste_numeric_cells':numeric,'preserved_zero_values':zero,'source_reference_links':references,'withheld_or_text_cells':withheld,'history_records':count,'preserved_negative_history_values':negative,'forest_records':len(unique),'crop_rows':len(crop['rows']),'checks':['source hashes','source cell identity','percent scaling','source years and codebook dates','city and country identity','missingness and anomalies','projection separation','complete IRP histories','native production units and flags']},indent=2))
