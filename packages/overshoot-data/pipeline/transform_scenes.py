"""Normalize acquired primary-source workbooks; no interpolation or synthetic rows.

Run with the Codex primary-runtime Python (openpyxl installed).
Input snapshots reside alongside this script. Output mass units: metric tonnes.
"""
import sys
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(sys.argv[1]) if len(sys.argv)>1 else Path(__file__).resolve().parent
RETRIEVED = "2026-09-19"

def values(name):
    return list(load_workbook(ROOT / name, read_only=True, data_only=True).active.values)

# Source files contain end-of-year anthropogenic mass in teratonnes (dry weight).
# Keep source year alignment, rather than Figure1.py's beginning-of-year shift.
human = {}
for name in ["anthropogenic_mass_2015.xlsx", "anthropogenic_mass_2037.xlsx"]:
    for r in values(name)[1:]:
        if isinstance(r[0], int):
            human[r[0]] = sum(r[1:7]) * 1e12  # Excludes accumulated waste column.
biomass = {int(r[1]): r[0] * 1e12 for r in values("biomass_dry.xlsx")[1:] if isinstance(r[1], int)}
uncertainty = {int(r[1]): r[0] for r in values("biomass_dry_uc.xlsx")[1:] if isinstance(r[1], int)}
stock_rows = []
for year in sorted(human.keys() & biomass.keys()):
    if year > 2025:
        continue
    r = {"year": year, "human": human[year], "biomass": biomass[year], "projected": year > 2015}
    if year in uncertainty:
        sd = biomass[year] * uncertainty[year] / 100
        r.update(low=biomass[year] - sd, high=biomass[year] + sd)
    stock_rows.append(r)

wb = load_workbook(ROOT / "waw3-country.xlsx", read_only=True, data_only=True)
rows = list(wb["Country dataset"].values)
headers = rows[1]
fate_labels = {
    "waste_treatment_open_dumpsite_percent": "Open dumpsite",
    "waste_treatment_controlled_landfill_percent": "Controlled landfill",
    "waste_treatment_sanitary_landfill_landfill_gas_system_percent": "Sanitary landfill / landfill gas system",
    "waste_treatment_landfill_unspecified_percent": "Landfill, unspecified",
    "waste_treatment_anaerobic_digestion_percent": "Anaerobic digestion",
    "waste_treatment_compost_percent": "Composting",
    "waste_treatment_recycling_percent": "Recycling",
    "waste_treatment_incineration_percent": "Incineration",
    "waste_treatment_mbt_percent": "Mechanical-biological treatment",
    "waste_treatment_rdf_percent": "Refuse-derived fuel",
    "waste_treatment_other_percent": "Other treatment",
    "waste_uncollected_percent": "Uncollected",
    "waste_treatment_unaccounted_for_percent": "Unaccounted for",
}
countries = []
for row_number, row in enumerate(rows[2:], start=3):
    if not isinstance(row[0], str) or len(row[0]) != 3:
        continue
    data = dict(zip(headers, row))
    fate = []
    for key, label in fate_labels.items():
        v = data.get(key)
        if isinstance(v, (float, int)):
            assert 0 <= v <= 1, (row[0], key, v)
            fate.append({"label": label, "share": v * 100, "source_field": key})
    countries.append({
        "country": row[0], "name": row[2],
        "tonnes": row[8] if isinstance(row[8], (int, float)) else None,
        "year": int(row[7]) if isinstance(row[7], (int, float)) else None,
        "source_id": "worldbank-waw3-2026",
        "source_row": row_number,
        "fate": fate,
        "fate_year_note": "Treatment fields have no per-field reference year in the flat workbook; the generation year must not be assigned to treatment.",
    })

scenes = {
    "version": "2026-09-19",
    "stock": {
        "rows": stock_rows, "source_id": "elhacham-2020", "unit": "metric tonnes dry mass",
        "method_note": "Original end-of-year anthropogenic mass, excluding waste; summed six material columns. Biomass uses source-provided years only, without smoothing or interpolation. Historical values are research estimates. Anthropogenic values after 2015 are the publication's extrapolations; biomass values after 2017 are extrapolations. Low/high, where present, are biomass +/- one standard deviation. Gaps between source years remain gaps; connecting chart segments are guides, not annual observations.",
        "crossing": {"year": 2020, "uncertainty_years": 6, "source_id": "elhacham-2020"},
    },
    "circularity": {
        "share": 6.9, "year": 2025, "referenceYear": "2021", "source_id": "circle-cgr-2025",
        "secondary_tonnes": 7.3e9, "processed_input_tonnes": 106.1e9,
        "method_note": "Published Circularity Metric: secondary material input / total primary + secondary material input. Includes recycling and downcycling. Input share, not the share of all waste recycled. Report 2025 v1.0, data 2021. 2026 edition focuses on economic value loss and continues to reference this 6.9% materials metric.",
    },
    "waste": {
        "global": {"tonnes": 2.56e9, "year": 2022, "source_id": "worldbank-waw3-2026", "modeled": True, "method_note": "Published rounded global 2022 baseline estimate; country data are drawn from different reference years and projected to 2022 in the source. This is not a sum of the country reported-year values below."},
        "countries": countries,
    },
    "unep_waste_fate": {
        "year": 2020, "source_id": "unep-gwmo-2024",
        "categories": [
            {"label": "Recycling", "share": 19, "tonnes": 404271000},
            {"label": "Waste-to-energy", "share": 13, "tonnes": 274800000},
            {"label": "Landfilling", "share": 30, "tonnes": 641256000},
            {"label": "Uncontrolled", "share": 38, "tonnes": 805644000},
        ],
        "method_note": "GWMO 2024 Figure 7, page 21. Original tonnes reported in thousands, multiplied by 1,000. Keep this baseline separate from the World Bank 2022 baseline. Uncontrolled combines dumping and open burning and cannot be separated from this figure. Displayed percentages are source-rounded.",
    },
    "material_system_2021": {
        "source_id": "circle-cgr-2025", "year": 2021,
        "extraction_tonnes": 98.8e9, "secondary_input_tonnes": 7.3e9, "processed_input_tonnes": 106.1e9,
        "gross_additions_to_stock_tonnes": 62.6e9, "demolition_and_discard_tonnes": 22.2e9,
        "net_additions_to_stock_tonnes": 40.4e9,
        "short_lived_materials_tonnes": 6.9e9, "energetic_use_tonnes": 36.6e9,
        "processed_output_tonnes": 65.7e9, "solid_and_liquid_waste_tonnes": 19.1e9,
        "emissions_to_air_tonnes": 43.5e9, "net_balancing_items_tonnes": 4.2e9,
        "method_note": "Direct transcription of CGR 2025 Figure 3 (page 25), original Gt multiplied by 1e9. The source Table 1 displays net stock 40.3 Gt while the Figure 3 flow diagram displays 40.4 Gt; retain diagram value for a diagram, document discrepancy, do not blend. Values rounded and not suitable for an exact mass-balance assertion. Waste output contains more than MSW. Never substitute current IRP extraction into these 2021 flow figures.",
    },
}

registry = [
    {
        "id": "elhacham-2020", "title": "Global human-made mass exceeds all living biomass", "publisher": "Elhacham et al.; Nature / Weizmann Institute of Science", "dataset": "Milo lab anthropogenic_mass source workbooks", "edition": "Nature 588, 442–444 (2020)", "publication_year": 2020, "coverage_years": [1900, 2037], "url": "https://www.nature.com/articles/s41586-020-3010-5", "data_url": "https://github.com/milo-lab/anthropogenic_mass", "citation": "Elhacham, E., Ben-Uri, L., Grozovski, J. et al. Global human-made mass exceeds all living biomass. Nature 588, 442–444 (2020). doi:10.1038/s41586-020-3010-5.", "license": "MIT for authors' data/code repository; Nature article has separate publisher rights.", "retrieved": RETRIEVED, "method_note": scenes["stock"]["method_note"], "unit_notes": "Source teratonnes dry weight, normalized by multiplying by 1e12.", "transformation_notes": "Six human-made categories summed; waste excluded; no smoothing/interpolation; joined only on available source years; shipped subset ends 2025.", "geographic_coverage": "Global", "temporal_coverage": "1900–2015 historical anthropogenic estimates; 2016–2037 source extrapolation; biomass estimates through 2017 with sparse years after 1990."
    },
    {
        "id": "circle-cgr-2025", "title": "The Circularity Gap Report 2025", "publisher": "Circle Economy, in collaboration with Deloitte", "dataset": "Circularity Indicator Set; Tables 1–2 and Figure 3", "edition": "Version 1.0, May 2025", "publication_year": 2025, "coverage_years": [2018, 2021], "url": "https://www.circle-economy.com/resources/the-circularity-gap-report-2025", "data_url": "https://circulareconomy.europa.eu/platform/sites/default/files/2025-09/CGR%202025%20complete%20document.pdf", "citation": "Circle Economy (2025). The circularity gap report 2025. Amsterdam: Circle Economy. Tables 1–2, pp. 24–25; Figure 3, p. 25.", "license": "CC BY-SA 4.0; the derived transcription is distributed under the same license with attribution.", "retrieved": RETRIEVED, "method_note": scenes["circularity"]["method_note"], "unit_notes": "Percent of total material input; flow figure original Gt normalized to metric tonnes.", "transformation_notes": "Exact source rounded numbers transcribed; no precision invented; percent not recalculated from rounded tonnes.", "geographic_coverage": "Global", "temporal_coverage": "2018 and 2021; shipped flow composition is 2021."
    },
    {
        "id": "worldbank-waw3-2026", "title": "What a Waste 3.0: Global Snapshot of Solid Waste Management Toward Circularity until 2050", "publisher": "World Bank Group", "dataset": "Country Dataset & Codebook", "edition": "March 2026", "publication_year": 2026, "coverage_years": [2010, 2050], "url": "https://www.worldbank.org/en/publication/what-a-waste", "data_url": "https://datacatalogfiles.worldbank.org/ddh-published/0039597/DR0095901/What_a_Waste_3.0_COUNTRY_Dataset_%26_Codebook.xlsx", "catalog_url": "https://datacatalog.worldbank.org/search/dataset/0039597/what-a-waste-global-database", "citation": "Cook, Ed, Ionkova, Kremena, Bhada-Tata, Perinaz, Yadav, Sonakshi, and Van Woerden, Frank. 2026. What a Waste 3.0: Global Snapshot of Solid Waste Management Toward Circularity until 2050. Country dataset (Excel). Washington, DC: World Bank Group.", "license": "CC BY 4.0", "retrieved": RETRIEVED, "method_note": "Country MSW values use the workbook's reported generation year, not projected 2022 column. Reporting definitions/measurement points differ. Fate fields have no per-field year in flat workbook and must not inherit generation year. Missing remains null/absent; shares are not renormalized.", "unit_notes": "Source MSW tonnes/year retained. Source numeric percentage cells are fractions; multiplied by 100 for percent display.", "transformation_notes": "Sheet Country dataset, row 2 machine field names, rows 3–219. Preserves ISO3, actual year and source row; extraction only, no estimation.", "geographic_coverage": "217 countries/economies", "temporal_coverage": "Country-specific reported years, projections to 2050 in source; projected country columns not used here."
    },
    {
        "id": "unep-gwmo-2024", "title": "Global Waste Management Outlook 2024: Beyond an Age of Waste – Turning Rubbish into a Resource", "publisher": "UNEP / ISWA", "dataset": "Figure 7: Global municipal solid waste destinations", "edition": "2024", "publication_year": 2024, "coverage_years": [2020, 2050], "url": "https://www.unep.org/resources/global-waste-management-outlook-2024", "data_url": "https://wedocs.unep.org/bitstreams/daa56f4d-2479-4e10-88c6-4d65da463299/download", "citation": "United Nations Environment Programme (2024). Global Waste Management Outlook 2024: Beyond an age of waste – Turning rubbish into a resource. Nairobi. Figure 7, p. 21.", "license": "Educational/non-profit reproduction with attribution; report prohibits commercial reproduction without permission. No imagery or report text reproduced in data artifact.", "retrieved": RETRIEVED, "method_note": scenes["unep_waste_fate"]["method_note"], "unit_notes": "Source Figure 7 thousand tonnes, multiplied by 1,000; rounded percent retained.", "transformation_notes": "Transcribed factual values only, no data interpolation or geographic substitution.", "geographic_coverage": "Global", "temporal_coverage": "2020 baseline; projections through 2050 not included."
    },
]

assert len(countries) == 217, len(countries)
assert len({r["country"] for r in countries}) == len(countries)
assert all(r["human"] > 0 and r["biomass"] > 0 for r in stock_rows)
(ROOT / "scenes.json").write_text(json.dumps(scenes, ensure_ascii=False, separators=(",", ":")))
(ROOT / "source-registry.json").write_text(json.dumps(registry, ensure_ascii=False, indent=2))
(ROOT / "snapshots.json").write_text(json.dumps({p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "retrieved": RETRIEVED} for p in ROOT.glob("*.xlsx")}, indent=2))
print(f"Wrote {len(stock_rows)} stock points and {len(countries)} country records")
