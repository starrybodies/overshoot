"""Build reproducible mine-to-plant coal receipts from final 2025 EIA files.

Inputs are downloaded separately from the official EIA-923, EIA-860 and MSHA
Mines datasets. Only shared identifiers establish a connection. No geography
or monthly quantities are interpolated.
"""
import collections
import csv
import io
import json
import pathlib
import sys
import zipfile

import openpyxl

root = pathlib.Path(__file__).resolve().parents[4]
source = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else root
output = pathlib.Path(__file__).resolve().parents[2] / 'public/data/v30'
output.mkdir(parents=True, exist_ok=True)

def num(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else 0

with zipfile.ZipFile(source / 'msha_mines.zip') as archive:
    reader = csv.DictReader(io.TextIOWrapper(archive.open('Mines.txt'), encoding='latin1'), delimiter='|')
    mines = {row['MINE_ID'].lstrip('0'): row for row in reader}

with zipfile.ZipFile(source / 'eia860_2025.zip') as archive:
    book = openpyxl.load_workbook(io.BytesIO(archive.read('2___Plant_Y2025.xlsx')), read_only=True, data_only=True)
    plants = {row[2]: row for row in book.active.iter_rows(min_row=3, values_only=True) if isinstance(row[2], int)}

with zipfile.ZipFile(source / 'eia923_2025.zip') as archive:
    book = openpyxl.load_workbook(io.BytesIO(archive.read('EIA923_Schedules_2_3_4_5_M_12_2025_Final.xlsx')), read_only=True, data_only=True)
    generation = collections.defaultdict(float)
    for row in book['Page 1 Generation and Fuel Data'].iter_rows(min_row=7, values_only=True):
        if row[14] in ('BIT', 'SUB', 'LIG', 'WC', 'RC') and isinstance(row[0], int):
            generation[row[0]] += num(row[95])
    grouped = collections.defaultdict(lambda: {'quantity': 0, 'months': set(), 'modes': set(), 'suppliers': set(), 'names': set()})
    all_rows = []
    total = linked = 0
    for row in book['Page 5 Fuel Receipts and Costs'].iter_rows(min_row=6, values_only=True):
        if row[8] != 'Coal' or not isinstance(row[2], int) or not num(row[15]):
            continue
        total += 1
        mine_id = str(row[12]).strip().lstrip('0')
        if not mine_id.isdigit():
            continue
        linked += 1
        key = (row[2], mine_id)
        item = grouped[key]
        item['quantity'] += row[15]
        item['months'].add(row[1])
        if row[25] and row[25] != '.': item['modes'].add(str(row[25]))
        if row[14] and row[14] != '.': item['suppliers'].add(str(row[14]))
        if row[13] and row[13] != '.': item['names'].add(str(row[13]))
        all_rows.append({'plantId': row[2], 'mineId': mine_id, 'month': row[1], 'shortTons': row[15], 'fuelCode': row[7], 'transportMode': row[25] or None})

plant_totals = collections.Counter()
for (plant_id, mine_id), value in grouped.items():
    if plant_id in plants and generation[plant_id] > 0: plant_totals[plant_id] += value['quantity']

sources = [
    {'id': 'eia923-2025', 'title': 'Power Plant Operations Report, Schedule 2 fuel receipts and Schedule 3 generation, final 2025', 'publisher': 'U.S. Energy Information Administration', 'url': 'https://www.eia.gov/electricity/data/eia923/index.php', 'publishedAt': '2026-09', 'reviewedAt': '2026-09-25', 'locator': '2025 final workbook, Page 5 and Page 1', 'reuse': 'U.S. government data'},
    {'id': 'eia860-2025', 'title': 'Annual Electric Generator Report, Plant 2025', 'publisher': 'U.S. Energy Information Administration', 'url': 'https://www.eia.gov/electricity/data/eia860/index.php', 'publishedAt': '2026-09', 'reviewedAt': '2026-09-25', 'locator': '2025 Plant workbook, Plant Code / Latitude / Longitude', 'reuse': 'U.S. government data'},
    {'id': 'msha-mines-2026', 'title': 'Mines Data Set', 'publisher': 'Mine Safety and Health Administration', 'url': 'https://arlweb.msha.gov/OpenGovernmentData/OGIMSHA.asp', 'publishedAt': '2026-09-25', 'reviewedAt': '2026-09-25', 'locator': 'Mines.txt, MINE_ID / LATITUDE / LONGITUDE', 'reuse': 'U.S. government data'},
]

def valid(lon, lat):
    try:
        pair = [round(float(lon), 6), round(float(lat), 6)]
        return pair if -180 <= pair[0] <= 180 and -90 <= pair[1] <= 90 and pair != [0, 0] else None
    except (ValueError, TypeError):
        return None

networks = []
for plant_id, total_quantity in plant_totals.most_common():
    plant = plants[plant_id]
    plant_point = valid(plant[10], plant[9])
    if not plant_point: continue
    pairs = sorted(((mine_id, value) for (pid, mine_id), value in grouped.items() if pid == plant_id), key=lambda x: -x[1]['quantity'])
    nodes = [{'id': f'plant-{plant_id}', 'name': str(plant[3]), 'country': 'USA', 'stage': 'manufacturing', 'coordinates': plant_point, 'locationNote': f'EIA-860 2025 power plant coordinates; {plant[6]}, {plant[8]}.', 'locationSource': 'eia860-2025'}]
    links = []
    for mine_id, item in pairs:
        mine = mines.get(mine_id)
        point = valid(mine.get('LONGITUDE'), mine.get('LATITUDE')) if mine else None
        mine_name = mine['CURRENT_MINE_NAME'] if mine else sorted(item['names'])[0] if item['names'] else f'MSHA {mine_id}'
        nodes.append({'id': f'mine-{mine_id}', 'name': mine_name, 'country': 'USA', 'stage': 'extraction', 'coordinates': point, 'locationNote': ('MSHA mine record point; the survey location may represent the mine entrance or operation, not a loading point.' if point else 'No validated mine point retained.'), 'locationSource': 'msha-mines-2026' if point else None})
        modes = sorted(item['modes'])
        links.append({'id': f'{mine_id}-{plant_id}', 'from': f'mine-{mine_id}', 'to': f'plant-{plant_id}', 'materialForm': 'Coal delivered to ' + str(plant[3]), 'mode': 'rail' if modes == ['RR'] else 'unspecified', 'claim': f'EIA-923 reports {item["quantity"]:,.0f} short tons received during {len(item["months"])} reported months of 2025 from MSHA mine {mine_id} at this plant. The reported primary transport code(s) are {", ".join(modes) if modes else "not specified"}.', 'sourceIds': ['eia923-2025', 'eia860-2025', 'msha-mines-2026'], 'evidence': 'government-reported', 'quantity': round(item['quantity']), 'quantityUnit': 'short tons received', 'months': sorted(item['months']), 'period': '2025 final', 'geometry': 'schematic' if point else 'unmapped', 'scope': 'Mine ID and plant ID are matched across government files. The straight map connection does not locate rail tracks, roads, intermediate handling, or a particular train.'})
    networks.append({'id': f'coal-{plant_id}', 'material': 'fuels', 'title': f'Coal into {plant[3]}', 'region': f'{plant[6]}, United States', 'summary': f'{total_quantity:,.0f} short tons of named-mine receipts in 2025 across {len(pairs)} mine connections. EIA reports {generation[plant_id]:,.0f} MWh of generation from coal at this plant in 2025; generation is a separate plant measure, not an allocation to these deliveries.', 'nodes': nodes, 'links': links, 'gaps': ['Receipt dates and fuel consumption are different measures; a delivered ton may be stored and burned later.', 'No rail track, transshipment point, or electricity destination is established by these records.', 'Rows with no valid MSHA mine ID are excluded from named-mine totals.'], 'generationMWh': round(generation[plant_id])})

catalog = {'release': 'v30', 'reviewedAt': '2026-09-25', 'method': f'Joined final EIA-923 2025 fuel receipts by MSHA mine ID and EIA plant ID to MSHA mines and EIA-860 plant points. {linked:,} of {total:,} 2025 coal receipt rows have a valid mine ID. Monthly quantities are summed for each named mine and plant; no missing month, coordinate, or route is interpolated. The display includes all generating plants with named-mine receipts, sorted by total identified-mine receipts; the monthly data download retains every identified-mine row.', 'geometry': 'Sourced mine and plant points; schematic straight connections between them. The actual transport corridor is unknown.', 'sources': sources, 'networks': networks, 'coverage': {'coalReceipts': total, 'identifiedMineReceipts': linked, 'identifiedMinePlantPairs': len(grouped), 'displayPlants': len(networks), 'monthlyRows': len(all_rows)}}
(output / 'coal-paths.json').write_text(json.dumps(catalog, ensure_ascii=False, separators=(',', ':')))
(output / 'coal-receipts-2025.json').write_text(json.dumps({'release': 'v30', 'source': 'eia923-2025', 'unit': 'short tons', 'rows': all_rows}, separators=(',', ':')))
print(catalog['coverage'])
print('mapped links', sum(l['geometry'] == 'schematic' for n in networks for l in n['links']))
